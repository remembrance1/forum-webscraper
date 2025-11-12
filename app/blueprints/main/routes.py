# app/blueprints/main/routes.py
from flask import render_template, request, redirect, url_for, send_file, flash, session, jsonify
from urllib.parse import urlparse
import uuid, threading, io, os, math, time
from datetime import datetime
from flask_login import login_required, current_user

from .tasks import run_scan_task, RUNS
from .parser_utils import render_results_html
from .fetch_utils import BACKENDS
from . import bp

from app.extensions import db
from app.models import Scan
from json import dumps

APP_TITLE = "Flask Forum Link Scraper"

# top of routes.py imports
import re, html
TAG_RE = re.compile(r"<[^>]+>")

def _clean(s: str) -> str:
    if not s:
        return ""
    # unescape any &lt;...&gt; etc.
    s = html.unescape(s)
    # drop tags
    s = TAG_RE.sub("", s)
    # collapse whitespace
    return " ".join(s.split())

@bp.route("/scraper", methods=["GET","POST"])
def scraper():
    if request.method == "GET":
        session.pop("run_id", None)
        session.pop("scan_saved", None)
        return render_template("sb_scraper.html", title=APP_TITLE, backends=BACKENDS)

    session.pop("run_id", None)
    url = (request.form.get("url") or "").strip()
    keyword = (request.form.get("keyword") or "").strip()
    sub_keyword = (request.form.get("sub_keyword") or "").strip()

    match_text = request.form.get("match_text") == "on"
    match_url = request.form.get("match_url") == "on"
    same_domain = request.form.get("same_domain") == "on"
    referer = (request.form.get("referer") or "").strip() or None
    cookies_raw = (request.form.get("cookies") or "").strip() or None
    backend = (request.form.get("backend") or os.environ.get("FETCH_BACKEND","auto")).strip().lower()
    if backend not in BACKENDS:
        backend = "auto"

    try:
        max_pages = int(request.form.get("max_pages") or 2)
    except ValueError:
        max_pages = 2
    max_pages = max(1, max_pages)

    try:
        pause_ms = int(request.form.get("pause_ms") or 400)
    except ValueError:
        pause_ms = 400
    pause_seconds = max(0, pause_ms) / 1000.0

    parsed = urlparse(url)
    if not parsed.scheme or not parsed.netloc:
        flash("Please provide a full URL including https://", "error")
        return redirect(url_for("main.scraper"))
    if not keyword:
        flash("Keyword cannot be empty.", "error")
        return redirect(url_for("main.scraper"))

    run_id = str(uuid.uuid4())
    RUNS[run_id] = {
        "results": [],
        "meta": {"source_url": url, "keyword": keyword, "sub_keyword": sub_keyword},
        "progress": {
            "status": "queued",
            "current": 0,
            "total": max_pages,
            "pages_scanned": 0,
            "links_seen": 0,
            "matches": 0,
            "eta_seconds": None,
            "message": "Queued",
        },
    }
    session["run_id"] = run_id

    t = threading.Thread(target=run_scan_task, kwargs=dict(
        run_id=run_id, url=url, keyword=keyword, sub_keyword=sub_keyword,
        match_text=match_text, match_url=match_url, same_domain=same_domain,
        referer=referer, cookies_raw=cookies_raw, backend=backend,
        pause_seconds=pause_seconds, max_pages=max_pages
    ), daemon=True)
    t.start()

    return redirect(url_for("main.results", page=1))

@bp.get("/results")
def results():
    run_id = session.get("run_id")
    data = RUNS.get(run_id) if run_id else None
    if not data:
        flash("No results to display. Please run a new scan.", "error")
        return redirect(url_for("main.scraper"))

    status = data.get("progress", {}).get("status", "done")
    # Get results safely (default empty list)
    matches = data.get("results") or []
    # Dedupe them
    matches = _dedupe_by_url(matches)
    # Persist back so pagination + exports use the deduped version
    RUNS[run_id]["results"] = matches
    # Total AFTER dedupe
    total = len(matches)
    meta = data.get("meta", {})
    per_page = 30

    # --- NEW: persist once when done and authenticated ---
    if status == "done" and current_user.is_authenticated and not session.get("scan_saved"):
        try:
            scan = Scan(
                user_id=current_user.id,
                source_url=(meta.get("source_url") or "")[:2048],
                keyword=(meta.get("keyword") or "")[:255],
                subkeyword=meta.get("sub_keyword"),
                logic=meta.get("logic"),
                num_matches=len(matches),
                results_json=dumps(matches)  # optional
            )
            db.session.add(scan)
            db.session.commit()
            session["scan_saved"] = True
        except Exception as e:
            # don't break the page if saving fails
            flash(f"Saved results partially, but could not store history: {e}", "warning")
    # --- end NEW ---

    # ... keep your pagination and render as-is
    per_page = 30

    try:
        page = int(request.args.get("page", 1))
    except ValueError:
        page = 1

    total_pages = max(1, math.ceil(total / per_page))
    page = max(1, min(page, total_pages))
    start = (page - 1) * per_page
    end = start + per_page

    return render_template("sb_scraper_results.html",
        title=APP_TITLE,
        source_url=meta.get("source_url"),
        keyword=meta.get("keyword"),
        sub_keyword=meta.get("sub_keyword"),
        matches=matches[start:end],
        page=page, total_pages=total_pages, per_page=per_page, total=total,
        start_index=start, run_id=run_id, status=status
    )

@bp.get("/export/html")
def export_html():
    # Prefer current run; fall back to session cache if present
    run_id = session.get("run_id")
    run = RUNS.get(run_id) if run_id else None

    if run:
        matches = _dedupe_by_url(run.get("results", []))
        meta = run.get("meta", {})
        source_url = meta.get("source_url", "")
        keyword = meta.get("keyword", "")
    else:
        data = session.get("last_results")
        if not data:
            flash("No results to export yet. Run a scan first.", "error")
            return redirect(url_for("main.scraper"))
        matches = _dedupe_by_url(data.get("matches", []))
        source_url = data.get("source_url", "")
        keyword = data.get("keyword", "")

    html_content = render_results_html(matches, source_url, keyword)
    buf = io.BytesIO(html_content.encode("utf-8"))
    return send_file(
        buf,
        mimetype="text/html",
        as_attachment=True,
        download_name=f"links_{int(time.time())}.html",
    )

@bp.get("/export/csv")
def export_csv():
    run_id = session.get("run_id")
    run = RUNS.get(run_id) if run_id else None
    if not run:
        flash("No results to export yet. Run a scan first.", "error")
        return redirect(url_for("main.scraper"))

    matches = _dedupe_by_url(run.get("results", []))

    import csv, io
    s = io.StringIO()
    w = csv.writer(s)
    w.writerow(["#", "Text", "URL"])
    for i, item in enumerate(matches, start=1):
        text, url = _coerce_item(item)   # _coerce_item already strips HTML from text
        w.writerow([i, text or url, url])

    mem = io.BytesIO(s.getvalue().encode("utf-8-sig"))
    return send_file(
        mem,
        mimetype="text/csv",
        as_attachment=True,
        download_name=f"links_{int(time.time())}.csv",
    )

@bp.get("/export/xlsx")
def export_xlsx():
    run_id = session.get("run_id")
    run = RUNS.get(run_id) if run_id else None
    if not run:
        flash("No results to export yet. Run a scan first.", "error")
        return redirect(url_for("main.scraper"))

    matches = _dedupe_by_url(run.get("results", []))

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Links"
    ws.append(["#", "Text", "URL"])
    for i, item in enumerate(matches, start=1):
        text, url = _coerce_item(item)   # _coerce_item already strips HTML from text
        ws.append([i, text or url, url])

    # simple column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return send_file(
        mem,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"links_{int(time.time())}.xlsx",
    )

@bp.get("/")
def landing():
    return render_template("landing.html", title="Welcome", current_year=datetime.utcnow().year)

@bp.get("/dashboard")
@login_required
def dashboard():
    from sqlalchemy import func
    # SCRAPER stats (real)
    q = Scan.query.filter_by(user_id=current_user.id)
    total_scans = q.count()
    total_matches = db.session.query(func.coalesce(func.sum(Scan.num_matches), 0))\
                              .filter(Scan.user_id == current_user.id)\
                              .scalar()
    last = q.order_by(Scan.created_at.desc()).first()

    stats_scraper = {
        "total_scans": total_scans,
        "total_matches": int(total_matches or 0),
        "last_scan": last.created_at.strftime("%Y-%m-%d %H:%M") if last else None,
    }
    recent_scans = q.order_by(Scan.created_at.desc()).limit(5).all()

    # CRAWLER stats (optional – safe placeholders if model not present)
    stats_crawler = {"total_crawls": 0, "pages_visited": 0, "last_crawl": None}
    recent_crawls = []
    try:
        from app.models import Crawl  # will fail until you add it
        q_crawler = Crawl.query.filter_by(user_id=current_user.id)
        stats_crawler["total_crawls"] = q_crawler.count()
        stats_crawler["pages_visited"] = int(db.session.query(
            func.coalesce(func.sum(Crawl.pages_crawled), 0)
        ).filter(Crawl.user_id == current_user.id).scalar() or 0)
        last_crawl = q_crawler.order_by(Crawl.created_at.desc()).first()
        stats_crawler["last_crawl"] = last_crawl.created_at.strftime("%Y-%m-%d %H:%M") if last_crawl else None
        recent_crawls = q_crawler.order_by(Crawl.created_at.desc()).limit(5).all()
    except Exception:
        # no Crawl model yet — keep placeholders
        pass

    return render_template(
        "dashboard.html",
        title="Dashboard",
        stats_scraper=stats_scraper,
        stats_crawler=stats_crawler,
        recent_scans=recent_scans,
        recent_crawls=recent_crawls,
    )

@bp.get("/history")
@login_required
def history():
    scans = (Scan.query
                  .filter_by(user_id=current_user.id)
                  .order_by(Scan.created_at.desc())
                  .limit(200)
                  .all())
    crawls = []
    try:
        from app.models import Crawl
        crawls = (Crawl.query
                        .filter_by(user_id=current_user.id)
                        .order_by(Crawl.created_at.desc())
                        .limit(200)
                        .all())
    except Exception:
        pass

    return render_template("history.html", title="History", scans=scans, crawls=crawls)


@bp.get("/history/<int:scan_id>")
@login_required
def scan_detail(scan_id):
    from json import loads
    scan = (Scan.query
                 .filter_by(id=scan_id, user_id=current_user.id)
                 .first_or_404())

    # decode results (your results_json stores a list of links/objects)
    results = []
    try:
        results = loads(scan.results_json or "[]")
    except Exception:
        pass

    # simple pagination
    page = max(1, int(request.args.get("page", 1)))
    per_page = 30
    start = (page - 1) * per_page
    end = start + per_page
    page_items = results[start:end]
    total_pages = (len(results) + per_page - 1) // per_page

    return render_template(
        "scan_detail.html",
        title="Scan Results",
        scan=scan,
        items=page_items,
        page=page,
        total_pages=total_pages,
    )

@bp.post("/history/clear")
@login_required
def clear_history():
    """Delete all saved scan rows for the current user."""
    try:
        # SQLAlchemy 2.0–style (works on 1.4+ too):
        from sqlalchemy import delete
        db.session.execute(
            delete(Scan).where(Scan.user_id == current_user.id)
        )
        db.session.commit()
        flash("All history deleted.", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Could not clear history: {e}", "error")
    return redirect(url_for("main.history"))

@bp.get("/progress/<run_id>")
def progress(run_id):
    data = RUNS.get(run_id)
    if not data:
        return jsonify({"status": "missing"}), 404
    prog = data.get("progress", {"current": 0, "total": 0, "status": "idle"})
    return jsonify(prog)

def _dedupe_by_url(items):
    seen = set()
    out = []
    for it in items:
        # Reuse your coercer to robustly extract the URL across dict/tuple/str
        _, url = _coerce_item(it)
        if url and url not in seen:
            seen.add(url)
            out.append(it)
    return out

def _coerce_item(item):
    """Return (text, url) from dict|tuple|str."""
    if isinstance(item, dict):
        url = item.get("url") or ""
        text = item.get("title") or item.get("text") or ""
        return (_clean(text) or url, url)
    elif isinstance(item, (list, tuple)):
        if len(item) >= 2:
            return (_clean(item[0] or ""), (item[1] or ""))
        elif len(item) == 1:
            x = item[0] or ""
            return (_clean(x), x)
    x = str(item)
    return (_clean(x), x)
